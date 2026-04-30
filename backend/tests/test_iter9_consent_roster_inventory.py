"""Iter9 backend regression: data_consent guards, team-roster template/import XLSX,
hotel CRUD (tier/stars), inventory public reads.
"""
import io
import os
import uuid
import requests
import pytest

def _resolve_backend_url():
    url = os.environ.get("REACT_APP_BACKEND_URL")
    if not url:
        try:
            with open("/app/frontend/.env") as f:
                for line in f:
                    if line.startswith("REACT_APP_BACKEND_URL="):
                        url = line.split("=", 1)[1].strip()
                        break
        except Exception:
            pass
    assert url, "REACT_APP_BACKEND_URL not set"
    return url.rstrip("/")


BASE_URL = _resolve_backend_url()


def _mongo_user(email: str):
    """Direct pymongo lookup (needed because backend exposes no admin /users endpoint)."""
    try:
        from pymongo import MongoClient
    except Exception:
        return None
    # Read mongo config from backend/.env
    mongo_url = "mongodb://localhost:27017"
    db_name = "test_database"
    try:
        with open("/app/backend/.env") as f:
            for line in f:
                if line.startswith("MONGO_URL="):
                    mongo_url = line.split("=", 1)[1].strip().strip('"').strip("'")
                elif line.startswith("DB_NAME="):
                    db_name = line.split("=", 1)[1].strip().strip('"').strip("'")
    except Exception:
        pass
    client = MongoClient(mongo_url, serverSelectionTimeoutMS=2000)
    return client[db_name].users.find_one({"email": email.lower()}, {"_id": 0})
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "admin@futuresoccercup.com"
ADMIN_PASSWORD = "FSCAdmin2025!"


# ---------- helpers ----------
def _new_session():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def admin_client():
    s = _new_session()
    r = s.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
    assert r.status_code == 200, f"admin login failed: {r.status_code} {r.text}"
    return s


@pytest.fixture(scope="module")
def team_client():
    """Register a brand-new team manager (data_consent=true) and keep the session."""
    s = _new_session()
    email = f"TEST_iter9_team_{uuid.uuid4().hex[:8]}@test.com"
    r = s.post(
        f"{API}/auth/register-team",
        json={
            "email": email,
            "password": "TeamPass2025!",
            "manager_name": "Iter9 Coach",
            "team_name": f"TEST_ITER9_{uuid.uuid4().hex[:6]}",
            "category": "Sub-12",
            "event_type": "premier_par",
            "data_consent": True,
        },
    )
    assert r.status_code == 200, f"register-team failed: {r.status_code} {r.text}"
    data = r.json()
    s.team_id = data["team_id"]
    s.email = email
    return s


@pytest.fixture(scope="module")
def family_client():
    s = _new_session()
    email = f"TEST_iter9_fam_{uuid.uuid4().hex[:8]}@test.com"
    r = s.post(
        f"{API}/auth/register",
        json={
            "email": email,
            "password": "FamPass2025!",
            "name": "Iter9 Familiar",
            "data_consent": True,
        },
    )
    assert r.status_code == 200
    return s


# ============================================================
# 1. Data consent guards
# ============================================================
class TestConsentGuards:
    def test_register_without_consent_rejected(self):
        r = requests.post(
            f"{API}/auth/register",
            json={
                "email": f"TEST_iter9_noconsent_{uuid.uuid4().hex[:6]}@test.com",
                "password": "X12345678!",
                "name": "Sin Consent",
                # no data_consent → default False
            },
        )
        assert r.status_code == 400
        assert "aceptar" in r.text.lower() or "consent" in r.text.lower() or "tratamiento" in r.text.lower()

    def test_register_team_without_consent_rejected(self):
        r = requests.post(
            f"{API}/auth/register-team",
            json={
                "email": f"TEST_iter9_tnc_{uuid.uuid4().hex[:6]}@test.com",
                "password": "X12345678!",
                "manager_name": "Coach NC",
                "team_name": f"TEST_NC_{uuid.uuid4().hex[:4]}",
                "category": "Sub-12",
                "event_type": "premier_par",
                "data_consent": False,
            },
        )
        assert r.status_code == 400
        assert "aceptar" in r.text.lower() or "consent" in r.text.lower() or "tratamiento" in r.text.lower()

    def test_register_with_consent_persists_fields(self):
        email = f"TEST_iter9_okfam_{uuid.uuid4().hex[:8]}@test.com"
        r = requests.post(
            f"{API}/auth/register",
            json={"email": email, "password": "OkPass2025!", "name": "Ok Fam", "data_consent": True},
        )
        assert r.status_code == 200
        doc = _mongo_user(email)
        assert doc is not None, f"user {email} not found in mongo"
        assert doc.get("data_consent") is True
        assert isinstance(doc.get("consent_at"), str) and len(doc["consent_at"]) > 10

    def test_register_team_with_consent_persists_fields(self, team_client):
        doc = _mongo_user(team_client.email)
        assert doc is not None
        assert doc.get("data_consent") is True
        assert isinstance(doc.get("consent_at"), str)


# ============================================================
# 2. team-roster template + import
# ============================================================
class TestTeamRoster:
    def test_template_forbidden_for_family(self, family_client):
        r = family_client.get(f"{API}/team-roster/template")
        assert r.status_code == 403

    def test_template_download_xlsx_as_team(self, team_client):
        r = team_client.get(f"{API}/team-roster/template")
        assert r.status_code == 200, r.text
        ct = r.headers.get("content-type", "")
        assert "spreadsheet" in ct or "officedocument" in ct
        content = r.content
        assert content[:2] == b"PK"  # xlsx = zip
        # validate 2 sheets
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not available in test env")
        wb = load_workbook(io.BytesIO(content))
        sn = [s.lower() for s in wb.sheetnames]
        assert any("jugador" in s for s in sn), f"missing Jugadores sheet in {wb.sheetnames}"
        assert any("tecnic" in s or "cuerpo" in s for s in sn), f"missing Cuerpo Tecnico sheet in {wb.sheetnames}"

    def test_import_preview_true_does_not_save(self, team_client, admin_client):
        # download template
        tpl = team_client.get(f"{API}/team-roster/template")
        assert tpl.status_code == 200
        xlsx_bytes = tpl.content
        # upload preview=true
        files = {"file": ("roster.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        # requests.Session has JSON content-type; override for multipart
        r = requests.post(
            f"{API}/team-roster/import?preview=true",
            files=files,
            cookies=team_client.cookies,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["saved"] is False
        assert "players" in data and "staff" in data
        assert data["players"]["total"] >= 1
        assert data["players"]["ok"] >= 1
        assert data["staff"]["total"] >= 1
        assert data["staff"]["ok"] >= 1
        # confirm DB has no new players yet for this team (admin must pass status=pendiente)
        players = admin_client.get(f"{API}/players?team_id={team_client.team_id}&status=pendiente").json()
        assert players == [] or len(players) == 0

    def test_import_preview_false_inserts_and_merges(self, team_client, admin_client):
        tpl = team_client.get(f"{API}/team-roster/template")
        assert tpl.status_code == 200
        files = {"file": ("roster.xlsx", tpl.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(
            f"{API}/team-roster/import?preview=false",
            files=files,
            cookies=team_client.cookies,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["saved"] is True
        assert data["players"]["ok"] >= 1
        assert data["staff"]["ok"] >= 1
        # verify persistence (team_client sees own roster regardless of status)
        players = team_client.get(f"{API}/players?team_id={team_client.team_id}").json()
        assert len(players) >= 1
        assert any(p["name"] == "Carlos Pérez" for p in players)
        # verify staff merged into team.cuerpo_tecnico
        teams = admin_client.get(f"{API}/teams").json()
        team = next((t for t in teams if t["id"] == team_client.team_id), None)
        assert team is not None
        staff = team.get("cuerpo_tecnico") or []
        assert any(s.get("name") == "Pedro Coach" for s in staff)


# ============================================================
# 3. Hotel / Transport / Tour inventory
# ============================================================
class TestInventory:
    def test_hotels_public_get(self):
        r = requests.get(f"{API}/hotels")
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_transports_public_get(self):
        r = requests.get(f"{API}/transports")
        assert r.status_code == 200

    def test_tours_public_get(self):
        r = requests.get(f"{API}/tours")
        assert r.status_code == 200

    def test_hotel_post_requires_admin(self):
        r = requests.post(
            f"{API}/hotels",
            json={"name": "anon", "description": "x", "price_per_night": 100},
        )
        assert r.status_code in (401, 403)

    def test_hotel_admin_crud_with_tier_stars(self, admin_client):
        payload = {
            "name": f"TEST_ITER9_Hotel_{uuid.uuid4().hex[:6]}",
            "description": "Hotel for iter9 test",
            "address": "Calle 1",
            "price_per_night": 250000,
            "tier": "gold",
            "stars": 4,
            "capacity": 3,
        }
        r = admin_client.post(f"{API}/hotels", json=payload)
        assert r.status_code in (200, 201), r.text
        created = r.json()
        hid = created["id"]
        assert created.get("tier") == "gold"
        assert created.get("stars") == 4

        # GET verify persistence
        got = admin_client.get(f"{API}/hotels/{hid}")
        assert got.status_code == 200
        assert got.json()["tier"] == "gold"
        assert got.json()["stars"] == 4

        # PUT update
        upd = dict(payload)
        upd["tier"] = "diamond"
        upd["stars"] = 5
        r2 = admin_client.put(f"{API}/hotels/{hid}", json=upd)
        assert r2.status_code == 200, r2.text
        assert r2.json()["tier"] == "diamond"

        # DELETE
        r3 = admin_client.delete(f"{API}/hotels/{hid}")
        assert r3.status_code in (200, 204)
        r4 = admin_client.get(f"{API}/hotels/{hid}")
        assert r4.status_code == 404


# ============================================================
# 4. Regression: quote calc COP 7.000.000 for Premier gold double 4x3 + transport
# ============================================================
class TestQuoteRegression:
    def test_quote_calc_premier_gold_double_4x3_transport(self):
        payload = {
            "event_type": "premier_par",
            "category": "Sub-12",
            "lodging_tier": "gold",
            "room_type": "double",
            "pax": 4,
            "nights": 3,
            "includes_transport": True,
            "includes_tour": False,
        }
        r = requests.post(f"{API}/quotes/calculate", json=payload)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["total_amount"] == 7000000, f"expected 7.000.000 COP, got {data.get('total_amount')}"


# ============================================================
# 5. Regression: bookings endpoints removed
# ============================================================
class TestBookingsRemoved:
    def test_bookings_list_404(self):
        r = requests.get(f"{API}/bookings")
        assert r.status_code in (404, 405)

    def test_bookings_mine_404(self):
        r = requests.get(f"{API}/bookings/mine")
        assert r.status_code in (404, 405, 401)
