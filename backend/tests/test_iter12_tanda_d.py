"""Tanda D — Datos y Estadisticas + tournament CRUD + intergrupos + historical standings + plantilla matches.

Cubre todos los casos backend descritos en iteration_12 review request.
"""
import os
import uuid
import io
import pytest
import requests
from pathlib import Path

def _load_base_url() -> str:
    env_path = Path("/app/frontend/.env")
    for line in env_path.read_text().splitlines():
        if line.startswith("REACT_APP_BACKEND_URL="):
            return line.split("=", 1)[1].strip().rstrip("/")
    raise RuntimeError("REACT_APP_BACKEND_URL not found")

BASE_URL = _load_base_url()
ADMIN_EMAIL = "admin@futuresoccercup.com"
ADMIN_PASSWORD = "FSCAdmin2025!"


@pytest.fixture(scope="module")
def admin():
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=20)
    assert r.status_code == 200, f"admin login failed: {r.text}"
    return s


@pytest.fixture(scope="module")
def dt():
    s = requests.Session()
    email = f"TEST_iter12_dt_{uuid.uuid4().hex[:6]}@test.com"
    pwd = "TeamPass2025!"
    r = s.post(f"{BASE_URL}/api/auth/register-team", json={
        "email": email, "password": pwd, "name": "DT Iter12",
        "team_name": f"TEST_ITER12_DT_{uuid.uuid4().hex[:6]}",
        "category": "Sub-12", "event_type": "festival", "data_consent": True,
    }, timeout=20)
    if r.status_code not in (200, 201):
        s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": pwd}, timeout=20)
    return s


# ===================== Tie-breaker: FP > DG =====================
class TestTieBreakerFP:
    GROUP = "TEST_FP_TBR"
    CATEGORY = "Sub-12"

    @pytest.fixture(scope="class")
    def teams(self, admin):
        ids = []
        for i in range(3):
            payload = {
                "name": f"TEST_FP_TEAM_{uuid.uuid4().hex[:6]}",
                "category": self.CATEGORY,
                "group_name": self.GROUP,
                "city": "Bogotá",
                "color": "#fff",
                "event_type": "festival",
            }
            r = admin.post(f"{BASE_URL}/api/teams", json=payload, timeout=20)
            assert r.status_code in (200, 201), r.text
            tid = r.json()["id"]
            # Approve
            admin.put(f"{BASE_URL}/api/teams/{tid}/status?status=aprobado", timeout=20)
            ids.append(tid)
        return ids

    def _create_match(self, admin, h, a):
        r = admin.post(f"{BASE_URL}/api/matches", json={
            "tournament_id": "",
            "home_team_id": h, "away_team_id": a,
            "match_date": "2030-07-01T10:00:00",
            "venue": "C1", "group_name": self.GROUP, "stage": "grupos",
        }, timeout=20)
        assert r.status_code in (200, 201), r.text
        return r.json()["id"]

    def _result(self, admin, mid, hs, as_, hfp, afp):
        r = admin.put(f"{BASE_URL}/api/matches/{mid}/result", json={
            "home_score": hs, "away_score": as_,
            "home_fair_play": hfp, "away_fair_play": afp,
        }, timeout=20)
        assert r.status_code == 200, r.text

    def test_fp_is_first_tiebreaker_above_gd(self, admin, teams):
        t1, t2, t3 = teams
        # Goal: equal points, but team t1 has highest FP and worst GD.
        # Make each team win exactly 1 match (3 pts each, 1 W 1 L 0 D).
        # t1 beats t2 1-0  (t1 FP +5, t2 FP +1)
        m1 = self._create_match(admin, t1, t2); self._result(admin, m1, 1, 0, 5, 1)
        # t2 beats t3 5-0  (t2 FP +1, t3 FP +1) — gives t2 best GD
        m2 = self._create_match(admin, t2, t3); self._result(admin, m2, 5, 0, 1, 1)
        # t3 beats t1 5-0  (t3 FP +1, t1 FP +1) — gives t1 worst GD
        m3 = self._create_match(admin, t3, t1); self._result(admin, m3, 5, 0, 1, 1)
        # Now: all 3 pts. t1 GD: 1-5=-4, FP=6. t2 GD: 5-1=+4, FP=2. t3 GD: 5-5=0, FP=2.
        r = admin.get(f"{BASE_URL}/api/stats/standings",
                      params={"category": self.CATEGORY, "group_name": self.GROUP}, timeout=20)
        assert r.status_code == 200, r.text
        rows = r.json()
        assert len(rows) == 3
        # t1 must be first because higher FP wins over worse GD
        assert rows[0]["team_id"] == t1, f"FP tiebreak failed. Got order: {[(x['team_id'], x['points'], x['fair_play'], x['gd']) for x in rows]}"
        assert rows[0]["fair_play"] >= rows[1]["fair_play"]


# ===================== Tournaments CRUD =====================
class TestTournaments:
    def test_get_archived_filter(self, admin):
        r_all = admin.get(f"{BASE_URL}/api/tournaments", timeout=20)
        assert r_all.status_code == 200
        all_t = r_all.json()
        # backfill defaults
        for t in all_t:
            assert "event_type" in t and "fmt" in t and "archived" in t

        r_arch = admin.get(f"{BASE_URL}/api/tournaments?archived=true", timeout=20)
        assert r_arch.status_code == 200
        for t in r_arch.json():
            assert t["archived"] is True
        r_act = admin.get(f"{BASE_URL}/api/tournaments?archived=false", timeout=20)
        assert r_act.status_code == 200
        for t in r_act.json():
            assert t["archived"] is False

    def test_post_extended_fields(self, admin):
        payload = {
            "name": f"TEST_TOUR_{uuid.uuid4().hex[:6]}",
            "season": "2030",
            "category": "Sub-12",
            "start_date": "2030-08-01",
            "end_date": "2030-08-30",
            "event_type": "festival",
            "fmt": "round_robin",
            "archived": False,
        }
        r = admin.post(f"{BASE_URL}/api/tournaments", json=payload, timeout=20)
        assert r.status_code in (200, 201), r.text
        data = r.json()
        assert data["event_type"] == "festival"
        assert data["fmt"] == "round_robin"
        assert data["archived"] is False
        assert "_id" not in data
        TestTournaments._tid = data["id"]

    def test_put_toggle_archived(self, admin):
        tid = getattr(TestTournaments, "_tid", None)
        assert tid
        r = admin.put(f"{BASE_URL}/api/tournaments/{tid}", json={"archived": True}, timeout=20)
        assert r.status_code == 200, r.text
        assert r.json()["archived"] is True
        assert "_id" not in r.json()

    def test_put_empty_body_400(self, admin):
        tid = getattr(TestTournaments, "_tid", None)
        r = admin.put(f"{BASE_URL}/api/tournaments/{tid}", json={}, timeout=20)
        assert r.status_code == 400

    def test_put_nonexistent_404(self, admin):
        r = admin.put(f"{BASE_URL}/api/tournaments/does-not-exist-xxxx", json={"name": "X"}, timeout=20)
        assert r.status_code == 404

    def test_put_no_auth_401_or_403(self, dt):
        # try without any session
        r = requests.put(f"{BASE_URL}/api/tournaments/whatever", json={"name": "X"}, timeout=20)
        assert r.status_code in (401, 403)
        # DT (non-admin)
        r2 = dt.put(f"{BASE_URL}/api/tournaments/whatever", json={"name": "X"}, timeout=20)
        assert r2.status_code in (401, 403)


# ===================== Intergroup =====================
class TestIntergroup:
    CATEGORY = "Sub-12"

    @pytest.fixture(scope="class")
    def setup_groups(self, admin):
        # Create 2 teams in group IG_A and 2 in group IG_B
        ids = {"IG_A": [], "IG_B": []}
        for grp in ["TEST_IG_A", "TEST_IG_B"]:
            for _ in range(2):
                r = admin.post(f"{BASE_URL}/api/teams", json={
                    "name": f"TEST_IG_TEAM_{uuid.uuid4().hex[:6]}",
                    "category": self.CATEGORY, "group_name": grp,
                    "city": "Bogotá", "color": "#fff", "event_type": "festival",
                }, timeout=20)
                assert r.status_code in (200, 201), r.text
                tid = r.json()["id"]
                admin.put(f"{BASE_URL}/api/teams/{tid}/status?status=aprobado", timeout=20)
                ids[grp.replace("TEST_", "")].append(tid)
        return ids

    def test_preview_no_persist(self, admin, setup_groups):
        body = {
            "tournament_id": "",
            "category": self.CATEGORY,
            "group_a": "TEST_IG_A",
            "group_b": "TEST_IG_B",
            "match_date": "2030-09-01",
            "pairing": "seed",
            "venues": ["C1", "C2"],
            "time_slots": ["10:00", "11:00"],
            "preview": True,
        }
        r = admin.post(f"{BASE_URL}/api/fixtures/intergroup", json=body, timeout=20)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["saved"] is False
        assert data["count"] == 2
        for m in data["matches"]:
            assert m["match_type"] == "intergrupo"
            assert "_id" not in m

    def test_save_persists(self, admin, setup_groups):
        body = {
            "tournament_id": "",
            "category": self.CATEGORY,
            "group_a": "TEST_IG_A",
            "group_b": "TEST_IG_B",
            "match_date": "2030-09-02",
            "pairing": "seed",
            "venues": ["C1"],
            "time_slots": ["10:00"],
            "preview": False,
        }
        r = admin.post(f"{BASE_URL}/api/fixtures/intergroup", json=body, timeout=20)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["saved"] is True
        assert data["count"] >= 1
        # Validate persistence in matches list
        r2 = admin.get(f"{BASE_URL}/api/matches", timeout=20)
        assert r2.status_code == 200
        matches = r2.json()
        intergrupos_on_date = [m for m in matches if m.get("match_type") == "intergrupo" and "2030-09-02" in (m.get("match_date") or "")]
        assert len(intergrupos_on_date) == data["count"]

    def test_invalid_group_400(self, admin):
        r = admin.post(f"{BASE_URL}/api/fixtures/intergroup", json={
            "tournament_id": "", "category": self.CATEGORY,
            "group_a": "NO_EXISTE_AAA", "group_b": "TEST_IG_B",
            "match_date": "2030-09-03", "preview": True,
        }, timeout=20)
        assert r.status_code == 400
        assert "equipos aprobados" in r.text.lower() or "no tiene" in r.text.lower()

    def test_invalid_date_400(self, admin, setup_groups):
        r = admin.post(f"{BASE_URL}/api/fixtures/intergroup", json={
            "tournament_id": "", "category": self.CATEGORY,
            "group_a": "TEST_IG_A", "group_b": "TEST_IG_B",
            "match_date": "not-a-date", "preview": True,
        }, timeout=20)
        assert r.status_code == 400
        assert "fecha" in r.text.lower()

    def test_no_auth(self, dt):
        r = requests.post(f"{BASE_URL}/api/fixtures/intergroup", json={
            "tournament_id": "", "category": "Sub-12",
            "group_a": "A", "group_b": "B", "match_date": "2030-09-01",
        }, timeout=20)
        assert r.status_code in (401, 403)
        r2 = dt.post(f"{BASE_URL}/api/fixtures/intergroup", json={
            "tournament_id": "", "category": "Sub-12",
            "group_a": "A", "group_b": "B", "match_date": "2030-09-01",
        }, timeout=20)
        assert r2.status_code in (401, 403)


# ===================== Historical standings =====================
class TestHistorical:
    def test_get_historical_diciembre_2025(self, admin):
        # First find the historic tournament
        r = admin.get(f"{BASE_URL}/api/tournaments?archived=true", timeout=20)
        assert r.status_code == 200
        archived = r.json()
        diciembre = [t for t in archived if "iciembre" in t.get("name", "") and "2025" in t.get("name", "")]
        assert diciembre, f"No se encontró torneo histórico Diciembre 2025. Archived: {[t['name'] for t in archived]}"
        tid = diciembre[0]["id"]
        TestHistorical._historic_tid = tid
        r2 = admin.get(f"{BASE_URL}/api/historical/standings?tournament_id={tid}", timeout=30)
        assert r2.status_code == 200
        rows = r2.json()
        assert len(rows) == 210, f"Esperaba 210 filas, got {len(rows)}"
        required = {"rank", "team_name", "played", "won", "drawn", "lost", "gf", "ga", "gd", "fair_play", "points"}
        first = rows[0]
        assert required.issubset(set(first.keys())), f"Missing keys. Got: {first.keys()}"
        for row in rows[:50]:
            assert "_id" not in row

    def test_post_bulk_create_and_delete(self, admin):
        fake_tid = f"TEST_TID_{uuid.uuid4().hex[:8]}"
        rows = [
            {"tournament_id": fake_tid, "category": "Sub-12", "group_name": "TEST_GR",
             "rank": 1, "team_name": "TEST_TEAM_X", "played": 3, "won": 3, "drawn": 0, "lost": 0,
             "gf": 9, "ga": 1, "gd": 8, "fair_play": 12, "points": 9},
            {"tournament_id": fake_tid, "category": "Sub-12", "group_name": "TEST_GR",
             "rank": 2, "team_name": "TEST_TEAM_Y", "played": 3, "won": 0, "drawn": 0, "lost": 3,
             "gf": 1, "ga": 9, "gd": -8, "fair_play": 8, "points": 0},
        ]
        r = admin.post(f"{BASE_URL}/api/historical/standings", json=rows, timeout=20)
        assert r.status_code == 200, r.text
        created = r.json()
        assert len(created) == 2
        for c in created:
            assert "_id" not in c
        r2 = admin.get(f"{BASE_URL}/api/historical/standings?tournament_id={fake_tid}", timeout=20)
        assert r2.status_code == 200
        assert len(r2.json()) == 2
        # delete
        rd = admin.delete(f"{BASE_URL}/api/historical/standings?tournament_id={fake_tid}", timeout=20)
        assert rd.status_code == 200
        assert rd.json()["deleted"] == 2

    def test_post_empty_returns_empty(self, admin):
        r = admin.post(f"{BASE_URL}/api/historical/standings", json=[], timeout=20)
        assert r.status_code == 200
        assert r.json() == []

    def test_post_no_auth(self, dt):
        r = requests.post(f"{BASE_URL}/api/historical/standings", json=[], timeout=20)
        assert r.status_code in (401, 403)
        r2 = dt.post(f"{BASE_URL}/api/historical/standings", json=[], timeout=20)
        assert r2.status_code in (401, 403)

    def test_delete_no_auth(self, dt):
        r = requests.delete(f"{BASE_URL}/api/historical/standings?tournament_id=x", timeout=20)
        assert r.status_code in (401, 403)


# ===================== Matches template XLSX =====================
class TestMatchesTemplate:
    def test_download_xlsx(self, admin):
        r = admin.get(f"{BASE_URL}/api/import/matches-template", timeout=20)
        assert r.status_code == 200, r.text
        ct = r.headers.get("content-type", "").lower()
        assert "openxmlformats" in ct or "spreadsheetml" in ct, f"unexpected content-type: {ct}"
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not available")
        wb = load_workbook(io.BytesIO(r.content))
        assert "Partidos" in wb.sheetnames, f"sheets: {wb.sheetnames}"
        ws = wb["Partidos"]
        headers = [c.value for c in ws[1]]
        expected = [
            "fecha (YYYY-MM-DD)", "hora (HH:MM)", "categoria", "grupo", "jornada", "fase",
            "match_type (regular|intergrupo)", "local", "visitante",
            "marcador_local", "marcador_visitante", "fair_play_local", "fair_play_visitante", "cancha",
        ]
        assert headers == expected, f"headers mismatch: {headers}"
        # At least 1 example row
        assert ws.max_row >= 2

    def test_no_auth(self, dt):
        r = requests.get(f"{BASE_URL}/api/import/matches-template", timeout=20)
        assert r.status_code in (401, 403)
        r2 = dt.get(f"{BASE_URL}/api/import/matches-template", timeout=20)
        assert r2.status_code in (401, 403)


# ===================== match_type field roundtrip =====================
class TestMatchTypeField:
    def test_post_intergrupo_match(self, admin):
        # Create 2 fresh teams
        ids = []
        for _ in range(2):
            r = admin.post(f"{BASE_URL}/api/teams", json={
                "name": f"TEST_MT_TEAM_{uuid.uuid4().hex[:6]}",
                "category": "Sub-12", "group_name": "TEST_MT_GR",
                "city": "Bogotá", "color": "#fff", "event_type": "festival",
            }, timeout=20)
            assert r.status_code in (200, 201), r.text
            tid = r.json()["id"]
            admin.put(f"{BASE_URL}/api/teams/{tid}/status?status=aprobado", timeout=20)
            ids.append(tid)

        r = admin.post(f"{BASE_URL}/api/matches", json={
            "tournament_id": "", "home_team_id": ids[0], "away_team_id": ids[1],
            "match_date": "2030-10-01T10:00:00", "venue": "C1", "group_name": "TEST_MT_GR",
            "stage": "grupos", "match_type": "intergrupo",
        }, timeout=20)
        assert r.status_code in (200, 201), r.text
        mid = r.json()["id"]
        assert r.json()["match_type"] == "intergrupo"

        # PUT toggle back to regular
        r2 = admin.put(f"{BASE_URL}/api/matches/{mid}", json={"match_type": "regular"}, timeout=20)
        assert r2.status_code == 200, r2.text
        # verify persistence
        r3 = admin.get(f"{BASE_URL}/api/matches", timeout=20)
        m = next((x for x in r3.json() if x["id"] == mid), None)
        assert m is not None
        assert m["match_type"] == "regular"
